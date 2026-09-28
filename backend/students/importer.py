"""Student-record import: parse → validate → preview → commit (docs/02 D3).

Rules:
* The college owns the file; the platform never invents records.
* A blank cell never erases an existing value (blank = "keep").
* Existing records are updated only in the fields that differ; the preview
  shows old → new for every change before anything is written.
* Rows with errors are reported and skipped; valid rows can still be committed.
* Account status (``status``) and the link to a user are never touched here.
"""

from __future__ import annotations

import csv
import hashlib
import hmac
import io
from datetime import date, datetime

import phonenumbers
from django.conf import settings
from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.validators import validate_email
from django.db import transaction
from django.utils import timezone
from django.utils.translation import gettext, gettext_lazy
from openpyxl import load_workbook

from audit.services import RequestMeta, record
from organization.models import Program

from .models import StudentImportBatch, StudentImportRow, StudentRecord

MAX_ROWS = 10_000
MAX_BYTES = 5 * 1024 * 1024

# How a column is named in messages (translated; the file may use either language).
COLUMN_LABELS = {
    "university_number": gettext_lazy("University number"),
    "full_name_ar": gettext_lazy("Name in Arabic"),
    "program_code": gettext_lazy("Program code"),
    "level": gettext_lazy("Level"),
    "course_code": gettext_lazy("Course code"),
}


def column_label(field: str) -> str:
    return str(COLUMN_LABELS.get(field, field))


# Accepted headers (Arabic or English) → field name.
HEADERS = {
    "university_number": "university_number",
    "الرقم الجامعي": "university_number",
    "full_name_ar": "full_name_ar",
    "الاسم": "full_name_ar",
    "الاسم بالعربية": "full_name_ar",
    "full_name_en": "full_name_en",
    "الاسم بالإنجليزية": "full_name_en",
    "program_code": "program_code",
    "program": "program_code",
    "رمز البرنامج": "program_code",
    "البرنامج": "program_code",
    "level": "level",
    "المستوى": "level",
    "email": "email",
    "البريد": "email",
    "البريد الإلكتروني": "email",
    "phone": "phone",
    "الهاتف": "phone",
    "gender": "gender",
    "الجنس": "gender",
    "birth_date": "birth_date",
    "تاريخ الميلاد": "birth_date",
    "national_id": "national_id",
    "الرقم الوطني": "national_id",
}
REQUIRED = ("university_number", "full_name_ar", "program_code", "level")
GENDERS = {
    "male": "male",
    "m": "male",
    "ذكر": "male",
    "female": "female",
    "f": "female",
    "أنثى": "female",
    "انثى": "female",
}
# Fields compared/updated on existing records (status and user are never imported).
UPDATABLE = (
    "full_name_ar",
    "full_name_en",
    "program_id",
    "level",
    "email",
    "phone_e164",
    "gender",
    "birth_date",
    "national_id_hash",
)


class ImportFileError(ValueError):
    """The file as a whole is unusable (format, size, headers)."""


def _hash_national_id(value: str) -> str:
    key = settings.SECRET_KEY.encode()
    return hmac.new(key, value.strip().encode(), hashlib.sha256).hexdigest()


def read_rows(name: str, content: bytes) -> list[dict[str, str]]:
    rows, _ = read_table(name, content, HEADERS, REQUIRED)
    return rows


def read_table(
    name: str, content: bytes, headers: dict[str, str], required: tuple[str, ...]
) -> tuple[list[dict], list[str]]:
    """Read the first sheet (xlsx) or a UTF-8 CSV into dicts keyed by field name.

    Returns the rows and the detected field columns. Blank lines are skipped.
    """
    if len(content) > MAX_BYTES:
        raise ImportFileError(gettext("The file is larger than 5 MB."))
    lower = name.lower()
    if lower.endswith(".csv"):
        text = content.decode("utf-8-sig")
        table = list(csv.reader(io.StringIO(text)))
    elif lower.endswith((".xlsx", ".xlsm")):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        table = [["" if c is None else c for c in row] for row in sheet.iter_rows(values_only=True)]
    else:
        raise ImportFileError(gettext("Upload an Excel (.xlsx) or CSV file."))
    if not table:
        raise ImportFileError(gettext("The file is empty."))
    header = [headers.get(str(h).strip().lower(), headers.get(str(h).strip())) for h in table[0]]
    missing = [f for f in required if f not in header]
    if missing:
        raise ImportFileError(
            gettext("Missing required columns: %(join)s.")
            % {"join": "، ".join(column_label(f) for f in missing)}
        )
    rows = []
    for values in table[1:]:
        if not any(str(v).strip() for v in values):
            continue  # skip blank lines
        row = {}
        for field, value in zip(header, values, strict=False):
            if field:
                row[field] = value if isinstance(value, (date, datetime)) else str(value).strip()
        rows.append(row)
    if len(rows) > MAX_ROWS:
        raise ImportFileError(
            gettext("The file has more than %(MAX_ROWS)s rows.") % {"MAX_ROWS": MAX_ROWS}
        )
    return rows, [h for h in header if h]


def _parse_date(value) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError


def _normalize(raw: dict, programs: dict[str, Program]) -> tuple[dict, list[str]]:
    errors: list[str] = []
    out: dict = {}
    for field in REQUIRED:
        if not str(raw.get(field, "")).strip():
            errors.append(gettext("%(field)s: required") % {"field": column_label(field)})
    number = str(raw.get("university_number", "")).strip()
    out["university_number"] = number
    out["full_name_ar"] = " ".join(str(raw.get("full_name_ar", "")).split())
    out["full_name_en"] = " ".join(str(raw.get("full_name_en", "")).split())

    program = programs.get(str(raw.get("program_code", "")).strip().upper())
    if raw.get("program_code") and program is None:
        errors.append(
            gettext("program_code: unknown program %(get)s")
            % {"get": repr(raw.get("program_code"))}
        )
    out["program_id"] = program.pk if program else None

    level_raw = str(raw.get("level", "")).strip()
    try:
        level = int(float(level_raw)) if level_raw else None
    except ValueError:
        level = None
        errors.append(gettext("level: must be a number"))
    if level is not None and program is not None and not 1 <= level <= program.levels_count:
        errors.append(
            gettext("level: must be between 1 and %(levels_count)s")
            % {"levels_count": program.levels_count}
        )
    out["level"] = level

    email = str(raw.get("email", "")).strip().lower()
    if email:
        try:
            validate_email(email)
        except DjangoValidationError:
            errors.append(gettext("email: invalid"))
    out["email"] = email

    phone = str(raw.get("phone", "")).strip()
    out["phone_e164"] = ""
    if phone:
        try:
            parsed = phonenumbers.parse(phone, "SD")
            if not phonenumbers.is_valid_number(parsed):
                raise phonenumbers.NumberParseException(0, "invalid")
            out["phone_e164"] = phonenumbers.format_number(
                parsed, phonenumbers.PhoneNumberFormat.E164
            )
        except phonenumbers.NumberParseException:
            errors.append(gettext("phone: invalid"))

    gender = str(raw.get("gender", "")).strip().lower()
    out["gender"] = GENDERS.get(gender, "") if gender else ""
    if gender and not out["gender"]:
        errors.append(gettext("gender: use male/female (ذكر/أنثى)"))

    birth = raw.get("birth_date", "")
    out["birth_date"] = None
    if birth:
        try:
            out["birth_date"] = _parse_date(birth)
        except ValueError:
            errors.append(gettext("birth_date: use YYYY-MM-DD or DD/MM/YYYY"))

    national_id = str(raw.get("national_id", "")).strip()
    out["national_id_hash"] = _hash_national_id(national_id) if national_id else ""
    return out, errors


def _diff(record_: StudentRecord, data: dict) -> dict:
    changes = {}
    for field in UPDATABLE:
        new = data.get(field)
        if new in ("", None):
            continue  # blank keeps the existing value
        old = getattr(record_, field)
        old_cmp = old.isoformat() if isinstance(old, date) else old
        if old_cmp != new:
            changes[field] = [
                "••••" if field == "national_id_hash" else old_cmp,
                "••••" if field == "national_id_hash" else new,
            ]
    return changes


def validate_file(meta: RequestMeta, name: str, uploaded) -> StudentImportBatch:
    content = uploaded.read()
    rows = read_rows(name, content)
    uploaded.seek(0)
    programs = {p.code.upper(): p for p in Program.objects.all()}
    existing = {
        r.university_number: r
        for r in StudentRecord.objects.filter(
            university_number__in=[str(r.get("university_number", "")).strip() for r in rows]
        )
    }
    seen: set[str] = set()
    counts = {"create": 0, "update": 0, "skip": 0, "error": 0}
    row_objects = []
    for index, raw in enumerate(rows, start=2):  # row 1 is the header
        normalized, errors = _normalize(raw, programs)
        number = normalized["university_number"]
        if number and number in seen:
            errors.append(gettext("university_number: duplicated in this file"))
        seen.add(number)
        changes = {}
        if errors:
            action = StudentImportRow.Action.ERROR
        elif number in existing:
            changes = _diff(existing[number], normalized)
            action = StudentImportRow.Action.UPDATE if changes else StudentImportRow.Action.SKIP
        else:
            action = StudentImportRow.Action.CREATE
        counts[action] += 1
        row_objects.append(
            StudentImportRow(
                row_no=index,
                raw={
                    k: (v.isoformat() if isinstance(v, (date, datetime)) else v)
                    for k, v in raw.items()
                    if k != "national_id"
                },
                normalized=normalized,
                action=action,
                changes=changes,
                errors=errors,
            )
        )
    with transaction.atomic():
        batch = StudentImportBatch.objects.create(
            file_name=name[:255],
            file=uploaded,
            uploaded_by=meta.actor,
            status=(
                StudentImportBatch.Status.HAS_ERRORS
                if counts["error"]
                else StudentImportBatch.Status.VALIDATED
            ),
            summary={"rows": len(rows), **counts},
        )
        for row in row_objects:
            row.batch = batch
        StudentImportRow.objects.bulk_create(row_objects)
        record(meta, "students.import_validate", batch, new=batch.summary)
    return batch


def commit(meta: RequestMeta, batch: StudentImportBatch) -> StudentImportBatch:
    """Apply create/update rows; error rows are skipped. One transaction."""
    if batch.status not in (
        StudentImportBatch.Status.VALIDATED,
        StudentImportBatch.Status.HAS_ERRORS,
    ):
        raise ValueError("This batch was already committed or rejected.")
    created = updated = 0
    with transaction.atomic():
        rows = batch.rows.exclude(
            action__in=[StudentImportRow.Action.ERROR, StudentImportRow.Action.SKIP]
        )
        numbers = [r.normalized["university_number"] for r in rows]
        existing = {
            r.university_number: r
            for r in StudentRecord.objects.select_for_update().filter(university_number__in=numbers)
        }
        programs = {p.pk: p for p in Program.objects.all()}
        for row in rows:
            data = row.normalized
            current = existing.get(data["university_number"])
            if current is None:
                StudentRecord.objects.create(
                    university_number=data["university_number"],
                    full_name_ar=data["full_name_ar"],
                    full_name_en=data["full_name_en"],
                    program=programs[data["program_id"]],
                    level=data["level"],
                    email=data["email"],
                    phone_e164=data["phone_e164"],
                    gender=data["gender"],
                    birth_date=data["birth_date"],
                    national_id_hash=data["national_id_hash"],
                )
                created += 1
            else:
                changes = _diff(current, data)
                for field in changes:
                    setattr(current, field, data[field])
                if changes:
                    current.save()
                    updated += 1
        batch.status = StudentImportBatch.Status.COMMITTED
        batch.committed_at = timezone.now()
        batch.committed_by = meta.actor
        batch.summary = {**batch.summary, "created": created, "updated": updated}
        batch.save(
            update_fields=["status", "committed_at", "committed_by", "summary", "updated_at"]
        )
        record(meta, "students.import_commit", batch, new=batch.summary)
    return batch


def reject(meta: RequestMeta, batch: StudentImportBatch) -> StudentImportBatch:
    if batch.status == StudentImportBatch.Status.COMMITTED:
        raise ValueError("A committed batch cannot be rejected.")
    batch.status = StudentImportBatch.Status.REJECTED
    batch.save(update_fields=["status", "updated_at"])
    record(meta, "students.import_reject", batch)
    return batch
